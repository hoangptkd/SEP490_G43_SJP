from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
import json
import re
import subprocess

ROOT = Path(r"C:\Users\HP\Desktop\FPT_EDU\SEP490\SEP490_G43_SJP")
SRC = Path(r"C:\Users\HP\Downloads\Report4_Software Design Specification (3).docx")
OUT = Path(r"C:\Users\HP\Downloads\Report4_Software Design Specification (3)_mermaid_class_diagrams_added.docx")
TMP_OUT = OUT.with_name(OUT.stem + "_tmp.docx")
MMD_DIR = Path(r"C:\Users\HP\Desktop\FPT_EDU\SEP490\diagram\class_diagram\report4_mermaid_missing")
IMG_DIR = MMD_DIR / "rendered"
NODE = Path(r"C:\Users\HP\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
NODE_MODULES = Path(r"C:\Users\HP\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\node_modules")

MMD_DIR.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)


def md(code: str) -> str:
    return "\n".join(line.rstrip() for line in code.strip().splitlines()) + "\n"


DIAGRAMS = {
    "email_verify": md("""
classDiagram
direction TB
class VerifyEmailPage {
  <<React Boundary>>
  +verifyToken(token) void
}
class AuthApiService {
  <<Frontend Service>>
  +verifyEmail(token) UserResponse
}
class AuthController {
  <<REST Controller>>
  -AuthService authService
  +verifyEmail(request) UserResponse
}
class AuthService {
  <<Service>>
  -EmailVerificationTokenRepository tokenRepository
  -UserRepository userRepository
  +verifyEmail(token) UserResponse
}
class EmailVerificationTokenRepository {
  <<Repository>>
  +findByToken(token) Optional~EmailVerificationToken~
  +save(token) EmailVerificationToken
}
class UserRepository {
  <<Repository>>
  +save(user) User
}
class EmailVerificationToken {
  <<Entity>>
  +String token
  +LocalDateTime expiresAt
  +LocalDateTime usedAt
}
class User {
  <<Entity>>
  +String email
  +String status
  +LocalDateTime emailVerifiedAt
}
class UserResponse {
  <<DTO>>
  +String id
  +String email
  +String role
  +String status
}
VerifyEmailPage ..> AuthApiService : verify email
AuthApiService ..> AuthController : POST /auth/verify-email
AuthController --> AuthService : delegates
AuthService --> EmailVerificationTokenRepository : validate token
AuthService --> UserRepository : activate user
EmailVerificationToken --> User : belongs to
AuthService ..> UserResponse : returns
"""),
    "forgot_reset_password": md("""
classDiagram
direction TB
class ForgotPasswordPage {
  <<React Boundary>>
  +requestReset(email) void
}
class ResetPasswordPage {
  <<React Boundary>>
  +resetPassword(token,password) void
}
class AuthApiService {
  <<Frontend Service>>
  +forgotPassword(email) void
  +resetPassword(token,password) void
}
class AuthController {
  <<REST Controller>>
  +forgotPassword(request) ResponseEntity
  +resetPassword(request) ResponseEntity
}
class AuthService {
  <<Service>>
  -UserRepository userRepository
  -PasswordEncoder passwordEncoder
  -EmailService emailService
  +createResetToken(email) void
  +resetPassword(token,password) void
}
class PasswordResetTokenRepository {
  <<Designed Repository>>
  +findByToken(token) Optional~PasswordResetToken~
  +save(token) PasswordResetToken
}
class UserRepository {
  <<Repository>>
  +findByEmail(email) Optional~User~
  +save(user) User
}
class PasswordResetToken {
  <<Designed Entity>>
  +String token
  +LocalDateTime expiresAt
  +LocalDateTime usedAt
}
class User {
  <<Entity>>
  +String email
  +String passwordHash
}
class EmailService {
  <<Service>>
  +sendResetPasswordEmail(email,link) void
}
ForgotPasswordPage ..> AuthApiService : request reset
ResetPasswordPage ..> AuthApiService : submit password
AuthApiService ..> AuthController : API calls
AuthController --> AuthService : delegates
AuthService --> UserRepository : find/update user
AuthService --> PasswordResetTokenRepository : create/validate token
AuthService --> EmailService : send link
PasswordResetToken --> User : for user
"""),
    "logout": md("""
classDiagram
direction TB
class HeaderLayout {
  <<React Boundary>>
  +logout() void
}
class AuthStorage {
  <<Frontend Util>>
  +clearAuthSession() void
  +getToken() String
}
class AuthApiService {
  <<Frontend Service>>
  +logout() void
}
class AuthController {
  <<REST Controller>>
  +logout() ResponseEntity
}
class JwtAuthenticationFilter {
  <<Security Filter>>
  -JwtUtil jwtUtil
  +doFilterInternal(...) void
}
class JwtUtil {
  <<Component>>
  +validateToken(token) Boolean
}
HeaderLayout ..> AuthStorage : clear local session
HeaderLayout ..> AuthApiService : POST logout
AuthApiService ..> AuthController : /auth/logout
JwtAuthenticationFilter --> JwtUtil : validate later requests
"""),
    "candidate_update_profile": md("""
classDiagram
direction TB
class CandidateProfilePage {
  <<React Boundary>>
  +loadProfile() void
  +saveProfile(form) void
}
class CandidateServiceFrontend {
  <<Frontend Service>>
  +getProfile() CandidateProfileResponse
  +updateProfile(request) CandidateProfileResponse
}
class CandidateController {
  <<REST Controller>>
  -CandidateService candidateService
  +getProfile() CandidateProfileResponse
  +updateProfile(request) CandidateProfileResponse
}
class CandidateService {
  <<Service>>
  -AuthService authService
  -CandidateProfileRepository profileRepository
  -CandidateSkillRepository skillRepository
  +getProfile() CandidateProfileResponse
  +updateProfile(request) CandidateProfileResponse
}
class AuthService {
  <<Service>>
  +getCurrentUser() User
}
class CandidateProfileRepository {
  <<Repository>>
  +findByUserId(id) Optional~CandidateProfile~
  +save(profile) CandidateProfile
}
class CandidateSkillRepository {
  <<Repository>>
  +deleteByCandidateId(id) void
  +save(skill) CandidateSkill
}
class CandidateProfile {
  <<Entity>>
  +String fullName
  +String phone
  +String location
  +List skills
}
class CandidateProfileRequest {
  <<DTO>>
  +String fullName
  +String phone
  +List skills
}
class CandidateProfileResponse {
  <<DTO>>
  +String id
  +String fullName
  +boolean applyReady
}
CandidateProfilePage ..> CandidateServiceFrontend : GET/PUT profile
CandidateServiceFrontend ..> CandidateController : API
CandidateController --> CandidateService : delegates
CandidateService --> AuthService : current user
CandidateService --> CandidateProfileRepository : load/save
CandidateService --> CandidateSkillRepository : sync skills
CandidateService ..> CandidateProfileResponse : returns
"""),
    "save_favorite_job": md("""
classDiagram
direction TB
class JobDetailPage {
  <<React Boundary>>
  +saveJob(jobId) void
  +unsaveJob(jobId) void
}
class CandidateController {
  <<REST Controller>>
  +saveJob(jobId) void
  +unsaveJob(jobId) void
  +savedJobs() List~JobResponse~
}
class CandidateService {
  <<Service>>
  -SavedJobRepository savedJobRepository
  -JobRepository jobRepository
  +saveJob(jobId) void
  +unsaveJob(jobId) void
  +getSavedJobs() List~JobResponse~
}
class SavedJobRepository {
  <<Repository>>
  +existsByCandidateIdAndJobId(...) boolean
  +findByCandidateIdAndJobId(...) Optional~SavedJob~
  +save(savedJob) SavedJob
}
class JobRepository {
  <<Repository>>
  +findById(id) Optional~Job~
}
class SavedJob {
  <<Entity>>
  +CandidateProfile candidate
  +Job job
  +LocalDateTime createdAt
}
class CandidateProfile {
  <<Entity>>
  +UUID id
}
class Job {
  <<Entity>>
  +UUID id
  +String title
  +String status
}
JobDetailPage ..> CandidateController : save/unsave
CandidateController --> CandidateService : delegates
CandidateService --> SavedJobRepository : check/save/delete
CandidateService --> JobRepository : validate job
SavedJob --> CandidateProfile : candidate
SavedJob --> Job : job
"""),
    "view_cv_list": md("""
classDiagram
direction TB
class CVPage {
  <<React Boundary>>
  +loadCvs() void
  +setDefaultCv(id) void
}
class CandidateController {
  <<REST Controller>>
  +getCvs() List~CvResponse~
  +setDefaultCv(id) CvResponse
}
class CandidateService {
  <<Service>>
  -CandidateCvRepository cvRepository
  +getCvs() List~CvResponse~
  +setDefaultCv(id) CvResponse
}
class CandidateCvRepository {
  <<Repository>>
  +findByCandidateIdOrderByCreatedAtDesc(id) List~CandidateCv~
  +findByIdAndCandidateId(id,candidateId) Optional~CandidateCv~
}
class CandidateProfile {
  <<Entity>>
  +UUID id
}
class CandidateCv {
  <<Entity>>
  +UUID id
  +String originalFileName
  +boolean defaultCv
  +boolean deleted
}
class CvResponse {
  <<DTO>>
  +String id
  +String originalFileName
  +boolean defaultCv
}
CVPage ..> CandidateController : GET cvs
CandidateController --> CandidateService : delegates
CandidateService --> CandidateCvRepository : query/update default
CandidateCv --> CandidateProfile : belongs to
CandidateService ..> CvResponse : returns
"""),
    "delete_cv": md("""
classDiagram
direction TB
class CVPage {
  <<React Boundary>>
  +deleteCv(id) void
}
class CandidateController {
  <<REST Controller>>
  +deleteCv(id) void
}
class CandidateService {
  <<Service>>
  -CandidateCvRepository cvRepository
  -ApplicationRepository applicationRepository
  +deleteCv(id) void
}
class CandidateCvRepository {
  <<Repository>>
  +findByIdAndCandidateId(id,candidateId) Optional~CandidateCv~
  +delete(cv) void
}
class ApplicationRepository {
  <<Repository>>
  +existsByCvId(cvId) boolean
}
class CandidateCv {
  <<Entity>>
  +UUID id
  +boolean deleted
  +boolean defaultCv
}
class Application {
  <<Entity>>
  +UUID id
  +ApplicationStatus status
}
CVPage ..> CandidateController : DELETE cv
CandidateController --> CandidateService : delegates
CandidateService --> CandidateCvRepository : load/delete
CandidateService --> ApplicationRepository : check used CV
Application --> CandidateCv : references
CandidateService --> CandidateCv : soft delete if used
"""),
    "candidate_payment_history": md("""
classDiagram
direction TB
class PaymentHistoryPage {
  <<Boundary>>
  +loadPaymentHistory() void
}
class PaymentController {
  <<Designed Controller>>
  +getPaymentHistory() List~Payment~
}
class CandidateController {
  <<REST Controller>>
  +subscription() SubscriptionResponse
}
class PaymentService {
  <<Designed Service>>
  +getPaymentHistory(userId) List~Payment~
}
class CandidateService {
  <<Service>>
  +getSubscription() SubscriptionResponse
}
class PaymentRepository {
  <<Designed Repository>>
  +findByUserIdOrderByCreatedAtDesc(id) List~Payment~
}
class SubscriptionRepository {
  <<Repository>>
  +findTopByUserIdOrderByStartedAtDesc(id) Optional~Subscription~
}
class Payment {
  <<Designed Entity>>
  +String sessionId
  +BigDecimal amount
  +String status
}
class Subscription {
  <<Entity>>
  +String status
  +LocalDateTime expiresAt
}
class Plan {
  <<Entity>>
  +String code
  +String name
}
PaymentHistoryPage ..> PaymentController : GET history
PaymentHistoryPage ..> CandidateController : GET subscription
PaymentController --> PaymentService : delegates
PaymentService --> PaymentRepository : query payments
CandidateController --> CandidateService : delegates
CandidateService --> SubscriptionRepository : current subscription
Subscription --> Plan : plan
"""),
    "cancel_subscription": md("""
classDiagram
direction TB
class SubscriptionPage {
  <<React Boundary>>
  +cancelSubscription() void
}
class SubscriptionController {
  <<Designed Controller>>
  +cancelSubscription(id) ResponseEntity
}
class SubscriptionService {
  <<Designed Service>>
  -SubscriptionRepository subscriptionRepository
  +cancelSubscription(userId) Subscription
}
class SubscriptionRepository {
  <<Repository>>
  +findActiveByUserId(id) Optional~Subscription~
  +save(subscription) Subscription
}
class Subscription {
  <<Entity>>
  +String status
  +LocalDateTime startedAt
  +LocalDateTime expiresAt
}
class Plan {
  <<Entity>>
  +String code
  +String name
}
SubscriptionPage ..> SubscriptionController : POST cancel
SubscriptionController --> SubscriptionService : delegates
SubscriptionService --> SubscriptionRepository : load/save active
SubscriptionService --> Subscription : set CANCELLED
Subscription --> Plan : plan
"""),
    "employer_send_email": md("""
classDiagram
direction TB
class ApplicationsPage {
  <<Boundary>>
  +sendEmail(candidateId,message) void
}
class EmailController {
  <<Designed Controller>>
  +sendCandidateEmail(request) ResponseEntity
}
class EmailService {
  <<Service>>
  +sendInterviewScheduleEmail(...) void
  +sendApplicationResultEmail(...) void
  +sendCustomEmail(...) void
}
class EmployerService {
  <<Service>>
  +getCurrentEmployerOrRegisterPlaceholder() Employer
}
class ApplicationRepository {
  <<Repository>>
  +findById(id) Optional~Application~
}
class Employer {
  <<Entity>>
  +UUID id
  +Company company
}
class Application {
  <<Entity>>
  +UUID id
  +CandidateProfile candidate
}
class Notification {
  <<Entity>>
  +String type
  +String message
}
ApplicationsPage ..> EmailController : POST email
EmailController --> EmployerService : check employer
EmailController --> ApplicationRepository : validate target
EmailController --> EmailService : send message
EmailService ..> Notification : optional notice
"""),
    "manage_employer_profile": md("""
classDiagram
direction TB
class EmployerProfilePage {
  <<Boundary>>
  +loadProfile() void
  +updateProfile(form) void
}
class EmployerController {
  <<REST Controller>>
  +getCompanyProfile() CompanyProfileResponse
  +updateCompanyProfile(request) CompanyProfileResponse
}
class EmployerService {
  <<Service>>
  -AuthService authService
  -EmployerRepository employerRepository
  -CompanyRepository companyRepository
  +getCurrentEmployerOrRegisterPlaceholder() Employer
  +getCompanyProfile() CompanyProfileResponse
  +updateCompanyProfile(request) CompanyProfileResponse
}
class AuthService {
  <<Service>>
  +getCurrentUser() User
}
class EmployerRepository {
  <<Repository>>
  +findByUserId(id) Optional~Employer~
  +save(employer) Employer
}
class CompanyRepository {
  <<Repository>>
  +save(company) Company
}
class User {
  <<Entity>>
  +String email
  +String role
}
class Employer {
  <<Entity>>
  +String position
  +boolean owner
}
class Company {
  <<Entity>>
  +String name
  +String verificationStatus
}
EmployerProfilePage ..> EmployerController : GET/PUT profile
EmployerController --> EmployerService : delegates
EmployerService --> AuthService : current user
EmployerService --> EmployerRepository : load employer
EmployerService --> CompanyRepository : save company
Employer --> User : account
Employer --> Company : company
"""),
    "renew_job_posting": md("""
classDiagram
direction TB
class JobListingPage {
  <<React Boundary>>
  +renewJob(jobId,newDeadline) void
}
class EmployerController {
  <<REST Controller>>
  +updateJob(id,request) JobResponse
}
class EmployerService {
  <<Service>>
  +updateJob(id,request) JobResponse
}
class JobService {
  <<Service>>
  +updateJobResponse(id,request) JobResponse
  +findById(id) Job
}
class JobRepository {
  <<Repository>>
  +findById(id) Optional~Job~
  +save(job) Job
}
class SubscriptionService {
  <<Designed Service>>
  +checkRenewQuota(employer) boolean
}
class Job {
  <<Entity>>
  +String status
  +LocalDate deadline
}
class Employer {
  <<Entity>>
  +UUID id
}
JobListingPage ..> EmployerController : PUT deadline
EmployerController --> EmployerService : delegates
EmployerService --> JobService : owner update
JobService --> JobRepository : load/save
JobService ..> SubscriptionService : quota check
Job --> Employer : posted by
"""),
    "export_candidate_application_list": md("""
classDiagram
direction TB
class ApplicationsPage {
  <<Boundary>>
  +exportApplications(jobId) void
}
class ExportController {
  <<Designed Controller>>
  +exportApplications(jobId) File
}
class ApplicationService {
  <<Service>>
  +getApplicationsByJob(jobId,employer) List~ApplicationResponse~
}
class ExportService {
  <<Designed Service>>
  +toCsv(applications) File
  +toExcel(applications) File
}
class ApplicationRepository {
  <<Repository>>
  +findByJobId(jobId) List~Application~
}
class JobRepository {
  <<Repository>>
  +findById(jobId) Optional~Job~
}
class Application {
  <<Entity>>
  +String status
  +LocalDateTime submittedAt
}
class CandidateProfile {
  <<Entity>>
  +String fullName
  +String phone
}
class CandidateCv {
  <<Entity>>
  +String originalFileName
}
ApplicationsPage ..> ExportController : GET export
ExportController --> ApplicationService : load applications
ApplicationService --> JobRepository : check owner
ApplicationService --> ApplicationRepository : query applications
Application --> CandidateProfile : candidate
Application --> CandidateCv : CV
ExportController --> ExportService : generate file
"""),
    "manage_company_profile": md("""
classDiagram
direction TB
class CompanyProfilePage {
  <<React Boundary>>
  +loadCompany() void
  +saveCompany(form) void
  +uploadLogo(file) void
}
class EmployerController {
  <<REST Controller>>
  +getCompanyProfile() CompanyProfileResponse
  +updateCompanyProfile(request) CompanyProfileResponse
  +uploadCompanyLogo(file) CompanyProfileResponse
}
class EmployerService {
  <<Service>>
  -CompanyRepository companyRepository
  -CompanyLocationRepository locationRepository
  -Cloudinary cloudinary
  +getCompanyProfile() CompanyProfileResponse
  +updateCompanyProfile(request) CompanyProfileResponse
  +uploadCompanyLogo(file) CompanyProfileResponse
}
class EmployerRepository {
  <<Repository>>
  +findByUserId(id) Optional~Employer~
}
class CompanyRepository {
  <<Repository>>
  +findByName(name) Optional~Company~
  +save(company) Company
}
class CompanyLocationRepository {
  <<Repository>>
  +findByCompanyId(id) List~CompanyLocation~
}
class Employer {
  <<Entity>>
  +boolean owner
}
class Company {
  <<Entity>>
  +String name
  +String taxCode
  +String logoUrl
}
class CompanyLocation {
  <<Entity>>
  +String branchName
  +boolean headquarter
}
CompanyProfilePage ..> EmployerController : company API
EmployerController --> EmployerService : delegates
EmployerService --> EmployerRepository : check owner
EmployerService --> CompanyRepository : save profile
EmployerService --> CompanyLocationRepository : headquarters
Employer --> Company : owns
Company --> CompanyLocation : locations
"""),
    "view_all_posted_jobs": md("""
classDiagram
direction TB
class JobListingPage {
  <<React Boundary>>
  +loadJobs() void
}
class EmployerController {
  <<REST Controller>>
  +getCompanyJobs() List~JobResponse~
}
class EmployerService {
  <<Service>>
  -JobRepository jobRepository
  +getCompanyJobs() List~JobResponse~
}
class JobRepository {
  <<Repository>>
  +findByCompanyIdOrderByCreatedAtDesc(id) List~Job~
}
class Employer {
  <<Entity>>
  +Company company
}
class Company {
  <<Entity>>
  +String name
}
class Job {
  <<Entity>>
  +String title
  +String status
  +LocalDate deadline
}
class JobResponse {
  <<DTO>>
  +String id
  +String title
  +String status
}
JobListingPage ..> EmployerController : GET /employer/jobs
EmployerController --> EmployerService : delegates
EmployerService --> JobRepository : query company jobs
Employer --> Company : company
Job --> Company : belongs to
EmployerService ..> JobResponse : maps results
"""),
    "close_reopen_job_listing": md("""
classDiagram
direction TB
class JobListingPage {
  <<React Boundary>>
  +closeJob(id) void
  +reopenJob(id) void
}
class EmployerController {
  <<REST Controller>>
  +updateJob(id,request) JobResponse
}
class EmployerService {
  <<Service>>
  +updateJob(id,request) JobResponse
}
class JobService {
  <<Service>>
  +findById(id) Job
  +updateJobResponse(id,request) JobResponse
}
class JobRepository {
  <<Repository>>
  +findById(id) Optional~Job~
  +save(job) Job
}
class Job {
  <<Entity>>
  +String status
  +LocalDate deadline
}
class Employer {
  <<Entity>>
  +UUID id
}
JobListingPage ..> EmployerController : PUT status
EmployerController --> EmployerService : delegates
EmployerService --> JobService : owner update
JobService --> JobRepository : load/save
JobService --> Job : set CLOSED/PUBLISHED
Job --> Employer : posted by
"""),
    "delete_job_listing": md("""
classDiagram
direction TB
class JobListingPage {
  <<React Boundary>>
  +deleteJob(id) void
}
class EmployerController {
  <<REST Controller>>
  +deleteJob(id) void
}
class EmployerService {
  <<Service>>
  +deleteJob(id) void
}
class JobService {
  <<Service>>
  -JobRepository jobRepository
  -JobSkillRepository jobSkillRepository
  +deleteJobForEmployer(id,employer) void
}
class ApplicationRepository {
  <<Repository>>
  +countActiveByJobId(jobId) long
}
class JobRepository {
  <<Repository>>
  +findById(id) Optional~Job~
  +delete(job) void
}
class Job {
  <<Entity>>
  +String status
}
class Application {
  <<Entity>>
  +String status
}
JobListingPage ..> EmployerController : DELETE job
EmployerController --> EmployerService : delegates
EmployerService --> JobService : delete owner job
JobService ..> ApplicationRepository : active check
JobService --> JobRepository : load/delete
Application --> Job : belongs to
"""),
    "track_interview_status": md("""
classDiagram
direction TB
class InterviewStatusPage {
  <<Boundary>>
  +loadInterviewStatus(jobId) void
}
class InterviewController {
  <<Designed Controller>>
  +getInterviewStatus(jobId) List~InterviewSession~
}
class InterviewService {
  <<Designed Service>>
  +getInterviewStatus(jobId,employer) List~InterviewSession~
}
class InterviewSessionRepository {
  <<Repository>>
  +findByJobId(jobId) List~InterviewSession~
}
class ApplicationRepository {
  <<Repository>>
  +findByJobId(jobId) List~Application~
}
class InterviewSession {
  <<Entity>>
  +String status
  +LocalDateTime scheduledAt
}
class Application {
  <<Entity>>
  +String status
}
class CandidateProfile {
  <<Entity>>
  +String fullName
}
InterviewStatusPage ..> InterviewController : GET status
InterviewController --> InterviewService : delegates
InterviewService --> InterviewSessionRepository : sessions
InterviewService --> ApplicationRepository : applications
InterviewSession --> Application : for application
Application --> CandidateProfile : candidate
"""),
    "employer_payment_history": md("""
classDiagram
direction TB
class EmployerPaymentPage {
  <<Boundary>>
  +loadPaymentHistory() void
}
class PaymentController {
  <<Designed Controller>>
  +getEmployerPayments() List~Payment~
}
class PaymentService {
  <<Designed Service>>
  +getEmployerPaymentHistory(employerId) List~Payment~
}
class EmployerService {
  <<Service>>
  +getCurrentEmployerOrRegisterPlaceholder() Employer
}
class PaymentRepository {
  <<Designed Repository>>
  +findByEmployerIdOrderByCreatedAtDesc(id) List~Payment~
}
class Employer {
  <<Entity>>
  +UUID id
  +Company company
}
class Payment {
  <<Designed Entity>>
  +BigDecimal amount
  +String status
  +String sessionId
}
class Subscription {
  <<Entity>>
  +String status
  +Plan plan
}
EmployerPaymentPage ..> PaymentController : GET history
PaymentController --> EmployerService : check employer
PaymentController --> PaymentService : delegates
PaymentService --> PaymentRepository : query payments
Payment --> Employer : payer
Payment --> Subscription : activates
"""),
    "admin_job_catalog": md("""
classDiagram
direction TB
class JobCatalogPage {
  <<Boundary>>
  +loadCategories() void
  +saveCategory() void
  +saveSkill() void
}
class CatalogController {
  <<Designed Controller>>
  +listCategories() List~Category~
  +upsertCategory() Category
  +upsertSkill() Skill
}
class CatalogService {
  <<Designed Service>>
  -CategoryRepository categoryRepository
  -SkillRepository skillRepository
  +listCategories() List~Category~
  +saveCategory() Category
  +saveSkill() Skill
}
class CategoryRepository {
  <<Repository>>
  +findAll() List~Category~
  +save(category) Category
}
class SkillRepository {
  <<Repository>>
  +findByNameIgnoreCase(name) Optional~Skill~
  +save(skill) Skill
}
class Category {
  <<Entity>>
  +String name
  +String slug
}
class Skill {
  <<Entity>>
  +String name
  +String slug
  +String category
}
JobCatalogPage ..> CatalogController : catalog API
CatalogController --> CatalogService : delegates
CatalogService --> CategoryRepository : categories
CatalogService --> SkillRepository : skills
Skill --> Category : category
"""),
    "admin_master_data": md("""
classDiagram
direction TB
class MasterDataPage {
  <<Boundary>>
  +loadMasterData() void
  +updateItem() void
}
class MasterDataController {
  <<Designed Controller>>
  +listMasterData(type) List
  +updateMasterData(type,id) Object
}
class MasterDataService {
  <<Designed Service>>
  +listPlans() List~Plan~
  +listCategories() List~Category~
  +updateLookup() Object
}
class PlanRepository {
  <<Repository>>
  +findAll() List~Plan~
  +save(plan) Plan
}
class CategoryRepository {
  <<Repository>>
  +findAll() List~Category~
  +save(category) Category
}
class SkillRepository {
  <<Repository>>
  +findAll() List~Skill~
  +save(skill) Skill
}
class Plan {
  <<Entity>>
  +String code
  +String name
  +BigDecimal price
}
class Category {
  <<Entity>>
  +String name
  +String slug
}
class Skill {
  <<Entity>>
  +String name
  +String category
}
MasterDataPage ..> MasterDataController : GET/PUT data
MasterDataController --> MasterDataService : delegates
MasterDataService --> PlanRepository : plans
MasterDataService --> CategoryRepository : categories
MasterDataService --> SkillRepository : skills
"""),
    "admin_cancel_subscription": md("""
classDiagram
direction TB
class AdminSubscriptionPage {
  <<Boundary>>
  +cancelSubscription(subscriptionId) void
}
class AdminSubscriptionController {
  <<Designed Controller>>
  +cancelSubscription(id,reason) ResponseEntity
}
class AdminSubscriptionService {
  <<Designed Service>>
  -SubscriptionRepository subscriptionRepository
  +cancelSubscription(id,admin,reason) Subscription
}
class AuthService {
  <<Service>>
  +getCurrentUserResponse() UserResponse
}
class SubscriptionRepository {
  <<Repository>>
  +findById(id) Optional~Subscription~
  +save(subscription) Subscription
}
class Subscription {
  <<Entity>>
  +String status
  +LocalDateTime expiresAt
}
class Plan {
  <<Entity>>
  +String code
  +String name
}
class User {
  <<Entity>>
  +String email
  +String role
}
AdminSubscriptionPage ..> AdminSubscriptionController : POST cancel
AdminSubscriptionController --> AuthService : check admin
AdminSubscriptionController --> AdminSubscriptionService : delegates
AdminSubscriptionService --> SubscriptionRepository : load/save
Subscription --> Plan : plan
Subscription --> User : owner
"""),
    "admin_restore_subscription": md("""
classDiagram
direction TB
class AdminSubscriptionPage {
  <<Boundary>>
  +activateSubscription(userId,planId) void
  +restoreSubscription(id) void
}
class AdminSubscriptionController {
  <<Designed Controller>>
  +activateSubscription(request) ResponseEntity
  +restoreSubscription(id) ResponseEntity
}
class AdminSubscriptionService {
  <<Designed Service>>
  +activateSubscription(userId,planId) Subscription
  +restoreSubscription(id) Subscription
}
class UserRepository {
  <<Repository>>
  +findById(id) Optional~User~
}
class PlanRepository {
  <<Repository>>
  +findById(id) Optional~Plan~
}
class SubscriptionRepository {
  <<Repository>>
  +findById(id) Optional~Subscription~
  +save(subscription) Subscription
}
class User {
  <<Entity>>
  +String email
  +String role
}
class Plan {
  <<Entity>>
  +String code
  +BigDecimal price
}
class Subscription {
  <<Entity>>
  +String status
  +LocalDateTime expiresAt
}
AdminSubscriptionPage ..> AdminSubscriptionController : activate/restore
AdminSubscriptionController --> AdminSubscriptionService : delegates
AdminSubscriptionService --> UserRepository : validate user
AdminSubscriptionService --> PlanRepository : select plan
AdminSubscriptionService --> SubscriptionRepository : save active
Subscription --> User : owner
Subscription --> Plan : plan
"""),
}

TARGETS = [
    ("2.2 Common - Email Verify", "email_verify"),
    ("2.3 Common - Forgot and Reset Password", "forgot_reset_password"),
    ("2.4 Common - Logout", "logout"),
    ("2.7 Candidate - Update Profile", "candidate_update_profile"),
    ("2.9 Candidate - Save Favorite Job", "save_favorite_job"),
    ("2.12 Candidate - View CV List", "view_cv_list"),
    ("2.13 Candidate - Delete CV", "delete_cv"),
    ("2.18 Candidate - View Payment History", "candidate_payment_history"),
    ("2.19 Candidate - Cancel Subscription", "cancel_subscription"),
    ("2.26 Employer - Send Email", "employer_send_email"),
    ("2.27 Employer - Manage Employer Profile", "manage_employer_profile"),
    ("2.28 Employer - Renew Job Posting", "renew_job_posting"),
    ("2.29 Employer - Export Candidate Application List", "export_candidate_application_list"),
    ("2.31 Employer - Manage Company Profile", "manage_company_profile"),
    ("2.32 Employer - View All Posted Job Listings", "view_all_posted_jobs"),
    ("2.33 Employer - Close/Reopen Job Listing", "close_reopen_job_listing"),
    ("2.34 Employer - Delete Job Listing", "delete_job_listing"),
    ("2.35 Employer - Track interview Status", "track_interview_status"),
    ("2.36 Employer - View Payment History", "employer_payment_history"),
    ("2.40 Admin - System Catalog Management - Job Catalog", "admin_job_catalog"),
    ("2.41 Admin - System Catalog Management - Master Data", "admin_master_data"),
    ("2.42 Admin - Subscription Actions", "admin_cancel_subscription"),
    ("2.43 Admin - Subscription Actions", "admin_restore_subscription"),
]


def write_mmd_and_render():
    items = []
    for slug, code in DIAGRAMS.items():
        (MMD_DIR / f"{slug}.mmd").write_text(code, encoding="utf-8")
        items.append({"slug": slug, "code": code})
    json_path = MMD_DIR / "diagrams.json"
    json_path.write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
    env = dict(**__import__("os").environ)
    env["NODE_PATH"] = str(NODE_MODULES)
    subprocess.run([str(NODE), str(ROOT / "scripts" / "render_mermaid_diagrams.js"), str(json_path), str(IMG_DIR)], check=True, env=env)


def clean(text):
    return " ".join(text.split())


def set_para_text(para, text):
    for run in list(para.runs):
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def insert_image_before(paragraph, image_path):
    pic_para = paragraph.insert_paragraph_before("")
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_para.paragraph_format.keep_together = True
    run = pic_para.add_run()
    with Image.open(image_path) as img:
        width_px, height_px = img.size
    max_width = 6.5
    max_height = 7.55
    width_in = max_width
    projected_height = width_in * (height_px / width_px)
    if projected_height > max_height:
        width_in = max_height * (width_px / height_px)
    run.add_picture(str(image_path), width=Inches(width_in))
    return pic_para


def insert_heading_before(paragraph, text):
    h = paragraph.insert_paragraph_before(text)
    h.style = "Heading 4"
    h.paragraph_format.keep_with_next = True
    return h


def section_matches(text, target):
    return clean(text).startswith(target)


def build_doc():
    if TMP_OUT.exists():
        TMP_OUT.unlink()
    doc = Document(SRC)
    inserted = []
    for section_title, slug in TARGETS:
        paras = doc.paragraphs
        h3_idx = next((i for i, p in enumerate(paras) if p.style.name == "Heading 3" and section_matches(p.text, section_title)), None)
        if h3_idx is None:
            print("WARN missing section", section_title)
            continue
        end = len(paras)
        for j in range(h3_idx + 1, len(paras)):
            if paras[j].style.name in ("Heading 2", "Heading 3") and clean(paras[j].text):
                end = j
                break
        class_idx = None
        seq_idx = None
        for j in range(h3_idx + 1, end):
            txt = clean(paras[j].text).lower()
            if paras[j].style.name == "Heading 4" and "class diagram" in txt and class_idx is None:
                class_idx = j
            if paras[j].style.name == "Heading 4" and "sequence diagram" in txt and seq_idx is None:
                seq_idx = j
        match = re.match(r"2\.(\d+)\s+", clean(paras[h3_idx].text))
        number = match.group(1) if match else ""
        class_title = f"2.{number}.1 Class Diagram"
        seq_title = f"2.{number}.2 Sequence Diagram"
        image_path = IMG_DIR / f"{slug}.png"
        if class_idx is None:
            if seq_idx is None:
                print("WARN no sequence heading", section_title)
                continue
            set_para_text(paras[seq_idx], seq_title)
            insert_heading_before(paras[seq_idx], class_title)
            insert_image_before(paras[seq_idx], image_path)
        else:
            set_para_text(paras[class_idx], class_title)
            paras[class_idx].paragraph_format.keep_with_next = True
            if seq_idx is not None:
                set_para_text(paras[seq_idx], seq_title)
                target_idx = class_idx + 1 if class_idx + 1 < len(paras) else seq_idx
                insert_image_before(paras[target_idx], image_path)
            else:
                target = paras[class_idx + 1] if class_idx + 1 < len(paras) else paras[class_idx]
                insert_image_before(target, image_path)
        inserted.append(slug)
    doc.save(TMP_OUT)
    TMP_OUT.replace(OUT)
    print("Inserted Mermaid diagrams:", len(inserted))
    print("Saved:", OUT)


def main():
    write_mmd_and_render()
    build_doc()


if __name__ == "__main__":
    main()
