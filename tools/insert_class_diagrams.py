from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image


SOURCE_DOCX = Path(r"C:\Users\HP\Downloads\Report4_Software Design Specification.docx")
IMAGE_DIR = Path(r"C:\Users\HP\Desktop\FPT_EDU\SEP490\diagram\class_diagram_rendered")
OUTPUT_DOCX = Path(r"C:\Users\HP\Downloads\Report4_Software Design Specification_with_class_diagrams.docx")


DESCRIPTIONS = {
    "admin_dashboard_statistics": "Minh họa các lớp và mối quan hệ phục vụ chức năng thống kê trên bảng điều khiển của quản trị viên.",
    "admin_manage_job_content": "Mô tả cấu trúc lớp cho nghiệp vụ quản trị nội dung tin tuyển dụng, bao gồm kiểm duyệt và cập nhật thông tin bài đăng.",
    "admin_manage_user": "Thể hiện các lớp liên quan đến việc quản lý tài khoản người dùng và các thao tác điều hành của quản trị viên.",
    "ai_candidate_ranking": "Mô tả thành phần xếp hạng ứng viên bằng AI dựa trên hồ sơ, yêu cầu công việc và kết quả đánh giá.",
    "ai_job_recommendation": "Minh họa các lớp hỗ trợ gợi ý việc làm bằng AI dựa trên thông tin ứng viên và đặc điểm tin tuyển dụng.",
    "create_job": "Trình bày cấu trúc lớp cho quy trình nhà tuyển dụng tạo mới tin tuyển dụng trong hệ thống.",
    "employer_view_applications": "Mô tả các lớp hỗ trợ nhà tuyển dụng xem, lọc và quản lý danh sách ứng tuyển.",
    "login_account": "Thể hiện các lớp xử lý đăng nhập bằng tài khoản hệ thống và luồng xác thực người dùng.",
    "login_google": "Mô tả cấu trúc lớp cho đăng nhập bằng Google và việc liên kết thông tin xác thực bên ngoài.",
    "main_workflow": "Tổng hợp các lớp chính và mối quan hệ trong luồng nghiệp vụ tổng quát của hệ thống.",
    "manage_job": "Mô tả các lớp phục vụ việc quản lý tin tuyển dụng sau khi được tạo, bao gồm cập nhật trạng thái và nội dung.",
    "mock_ai_interview": "Minh họa cấu trúc lớp cho chức năng phỏng vấn thử bằng AI, câu hỏi, câu trả lời và kết quả đánh giá.",
    "payment_subscription": "Thể hiện các lớp liên quan đến gói dịch vụ, thanh toán và trạng thái đăng ký của người dùng.",
    "register_account": "Mô tả các lớp và quan hệ trong quy trình đăng ký tài khoản mới.",
    "schedule_interview_update_status": "Mô tả các lớp hỗ trợ đặt lịch phỏng vấn và cập nhật trạng thái phỏng vấn hoặc ứng tuyển.",
    "search_job": "Minh họa cấu trúc lớp cho chức năng tìm kiếm và lọc việc làm theo tiêu chí của ứng viên.",
    "setup_candidate": "Mô tả các lớp phục vụ việc thiết lập hồ sơ ứng viên sau khi tạo tài khoản.",
    "setup_employer": "Mô tả các lớp phục vụ việc thiết lập thông tin nhà tuyển dụng và công ty.",
    "submit_job_application": "Thể hiện các lớp liên quan đến quá trình ứng viên nộp đơn ứng tuyển vào một tin tuyển dụng.",
    "update_application_status": "Mô tả cấu trúc lớp cho việc cập nhật và theo dõi trạng thái hồ sơ ứng tuyển.",
    "upload_cv": "Minh họa các lớp xử lý việc tải lên CV, lưu trữ tệp và liên kết CV với hồ sơ ứng viên.",
    "view_interview_history_progress": "Mô tả các lớp hiển thị lịch sử phỏng vấn và tiến độ đánh giá của ứng viên.",
}


def clean_key(image_path: Path) -> str:
    suffix = "_class_diagram_vertical_fit"
    name = image_path.stem
    return name[: -len(suffix)] if name.endswith(suffix) else name


def title_from_key(key: str) -> str:
    special = {"ai": "AI", "cv": "CV"}
    return " ".join(special.get(word, word.capitalize()) for word in key.split("_"))


def fit_size(image_path: Path, max_width=6.5, max_height=7.0):
    with Image.open(image_path) as im:
        width_px, height_px = im.size
    ratio = width_px / height_px
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return Inches(width), Inches(height)


def add_caption(doc: Document, figure_no: int, title: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"Figure {figure_no}. {title} Class Diagram")
    run.bold = True
    run.font.size = Pt(10)


def add_description(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(10)


def main():
    copy2(SOURCE_DOCX, OUTPUT_DOCX)
    doc = Document(OUTPUT_DOCX)

    doc.add_page_break()
    doc.add_heading("Class Diagrams", level=1)
    intro = doc.add_paragraph(
        "This section presents the class diagrams converted from the Mermaid source files. "
        "Each diagram summarizes the main classes, responsibilities, and relationships for a key system function."
    )
    intro.paragraph_format.space_after = Pt(12)

    image_paths = sorted(IMAGE_DIR.glob("*.png"))
    for index, image_path in enumerate(image_paths, start=1):
        if index > 1:
            doc.add_page_break()
        key = clean_key(image_path)
        title = title_from_key(key)
        doc.add_heading(title, level=2)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width, height = fit_size(image_path)
        paragraph.add_run().add_picture(str(image_path), width=width, height=height)
        add_caption(doc, index, title)
        add_description(doc, DESCRIPTIONS.get(key, f"This diagram describes the class structure for {title.lower()}."))

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
