from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from shutil import copy2
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


SOURCE_DOCX = Path(r"C:\Users\HP\Downloads\Report4_Software Design Specification.docx")
OUTPUT_DOCX = Path(r"C:\Users\HP\Downloads\Report4_Software Design Specification_completed.docx")
CLASS_IMAGE_DIR = Path(r"C:\Users\HP\Desktop\FPT_EDU\SEP490\diagram\class_diagram_rendered")
CLASS_MMD_DIR = Path(r"C:\Users\HP\Desktop\FPT_EDU\SEP490\diagram\class_diagram")
SEQUENCE_IMAGE_DIR = Path(r"C:\Users\HP\Desktop\FPT_EDU\SEP490\usdd_run_jar\image")


@dataclass(frozen=True)
class Feature:
    key: str
    title: str
    description: str
    class_image: str
    sequence_image: str | None
    sequence_title: str
    sequence_description: str


FEATURES: list[Feature] = [
    Feature(
        "main_workflow",
        "Main Workflow Overview",
        "This overview describes the core design structure connecting authentication, candidate operations, employer operations, application processing, AI services, payment, and administration.",
        "main_workflow_class_diagram_vertical_fit.png",
        None,
        "Main Workflow Reference",
        "The end-to-end behavior is decomposed into the feature-level sequence diagrams in the following subsections.",
    ),
    Feature(
        "register_account",
        "Register Account",
        "This feature allows a guest user to create a new account and initialize the role-specific access path required by the recruitment platform.",
        "register_account_class_diagram_vertical_fit.png",
        "login-đăng ký.drawio.png",
        "Register Account Sequence Diagram",
        "The sequence shows the registration flow from user input validation through account creation, persistence, and response handling.",
    ),
    Feature(
        "login_account",
        "Login Account",
        "This feature authenticates users with system credentials, validates account status, and returns the session information required by the frontend.",
        "login_account_class_diagram_vertical_fit.png",
        "login-loginAccount.drawio.png",
        "Login Account Sequence Diagram",
        "The sequence presents credential submission, backend validation, password checking, token creation, and frontend session storage.",
    ),
    Feature(
        "login_google",
        "Login Google",
        "This feature supports Google-based authentication for users who choose an external identity provider instead of a local password.",
        "login_google_class_diagram_vertical_fit.png",
        "login-LoginGG.drawio.png",
        "Google Login Sequence Diagram",
        "The sequence describes Google OAuth handoff, token verification, local user lookup or creation, and session response generation.",
    ),
    Feature(
        "setup_candidate",
        "Setup Candidate",
        "This feature captures the candidate profile data required for job matching, application submission, and AI-assisted recommendations.",
        "setup_candidate_class_diagram_vertical_fit.png",
        "login-setup_candidate.drawio.png",
        "Setup Candidate Sequence Diagram",
        "The sequence shows profile form submission, candidate data validation, service processing, and profile persistence.",
    ),
    Feature(
        "upload_cv",
        "Upload CV",
        "This feature enables candidates to upload CV files, store file metadata, and link the uploaded CV to their candidate profile.",
        "upload_cv_class_diagram_vertical_fit.png",
        "login-uploadCV.drawio.png",
        "Upload CV Sequence Diagram",
        "The sequence covers file selection, upload request processing, storage handling, metadata saving, and response feedback.",
    ),
    Feature(
        "search_job",
        "Search Job",
        "This feature lets candidates browse and filter available job postings using criteria such as keyword, skills, location, and company.",
        "search_job_class_diagram_vertical_fit.png",
        "login-Seach job.drawio.png",
        "Search Job Sequence Diagram",
        "The sequence explains how search criteria are submitted, translated into repository queries, and returned as job result data.",
    ),
    Feature(
        "submit_job_application",
        "Submit Job Application",
        "This feature supports candidate application submission by combining the selected job, candidate profile, CV, and application status.",
        "submit_job_application_class_diagram_vertical_fit.png",
        "login-subbmit job.drawio.png",
        "Submit Job Application Sequence Diagram",
        "The sequence shows job selection, duplicate validation, application creation, persistence, and user confirmation.",
    ),
    Feature(
        "ai_job_recommendation",
        "AI Job Recommendation",
        "This feature recommends jobs to candidates by analyzing profile and CV information against available job posting requirements.",
        "ai_job_recommendation_class_diagram_vertical_fit.png",
        "login-Ai_gợi ý.drawio.png",
        "AI Job Recommendation Sequence Diagram",
        "The sequence describes profile retrieval, AI matching request preparation, recommendation scoring, and result presentation.",
    ),
    Feature(
        "setup_employer",
        "Setup Employer",
        "This feature collects employer and company information so employer accounts can publish jobs and manage applicants.",
        "setup_employer_class_diagram_vertical_fit.png",
        "login-setup-employer.drawio.png",
        "Setup Employer Sequence Diagram",
        "The sequence shows company profile input, employer validation, company persistence, and completion feedback.",
    ),
    Feature(
        "create_job",
        "Create Job",
        "This feature allows an employer to create a job posting and optionally use AI assistance to generate or refine job content.",
        "create_job_class_diagram_vertical_fit.png",
        "login-create job.drawio.png",
        "Create Job Sequence Diagram",
        "The sequence presents quota checking, job draft generation, job submission, validation, and storage of the posting.",
    ),
    Feature(
        "manage_job",
        "Manage Job",
        "This feature enables employers to update, publish, close, or review the status of their job postings.",
        "manage_job_class_diagram_vertical_fit.png",
        "login-Manage job.drawio.png",
        "Manage Job Sequence Diagram",
        "The sequence explains how an employer retrieves job data, performs job management actions, and receives updated results.",
    ),
    Feature(
        "employer_view_applications",
        "Employer View Applications",
        "This feature helps employers view applicant lists and inspect application information for their job postings.",
        "employer_view_applications_class_diagram_vertical_fit.png",
        "login-emplyer view application.drawio.png",
        "Employer View Applications Sequence Diagram",
        "The sequence describes retrieving applications by job or employer, loading related candidate data, and presenting the applicant list.",
    ),
    Feature(
        "update_application_status",
        "Update Application Status",
        "This feature allows employers to update application progress, such as shortlisted, rejected, interview, or hired states.",
        "update_application_status_class_diagram_vertical_fit.png",
        "login-update_application_status.drawio.png",
        "Update Application Status Sequence Diagram",
        "The sequence shows status-change request validation, application update, notification handling, and response delivery.",
    ),
    Feature(
        "ai_candidate_ranking",
        "AI Candidate Ranking",
        "This feature ranks candidates for a job by using AI-supported scoring based on CV content, profile data, and job requirements.",
        "ai_candidate_ranking_class_diagram_vertical_fit.png",
        "login-ai-candidate-ranking.drawio.png",
        "AI Candidate Ranking Sequence Diagram",
        "The sequence covers employer ranking request, candidate data collection, AI scoring, ranking persistence, and display.",
    ),
    Feature(
        "schedule_interview_update_status",
        "Schedule Interview Update Status",
        "This feature supports interview scheduling and keeps application or interview status synchronized with the selected schedule.",
        "schedule_interview_update_status_class_diagram_vertical_fit.png",
        "login-hẹn phỏng vấn + đổi trạng thái.drawio.png",
        "Schedule Interview and Update Status Sequence Diagram",
        "The sequence shows schedule creation, candidate notification, interview status update, and application status alignment.",
    ),
    Feature(
        "mock_ai_interview",
        "Mock AI Interview",
        "This feature allows candidates to practice interviews with AI-generated questions and receive evaluation feedback.",
        "mock_ai_interview_class_diagram_vertical_fit.png",
        "login-phỏng vấn ảo.drawio.png",
        "Mock AI Interview Sequence Diagram",
        "The sequence describes starting a practice session, generating questions, recording answers, evaluating responses, and saving feedback.",
    ),
    Feature(
        "view_interview_history_progress",
        "View Interview History Progress",
        "This feature lets candidates view historical interview sessions and track progress across previous AI interview evaluations.",
        "view_interview_history_progress_class_diagram_vertical_fit.png",
        "login-view_interview_history_progress.drawio.png",
        "View Interview History and Progress Sequence Diagram",
        "The sequence shows retrieval of interview history, aggregation of progress metrics, and display of candidate performance details.",
    ),
    Feature(
        "payment_subscription",
        "Payment Subscription",
        "This feature manages subscription plan selection, payment processing, entitlement updates, and subscription status tracking.",
        "payment_subscription_class_diagram_vertical_fit.png",
        "login-payment.drawio.png",
        "Payment Subscription Sequence Diagram",
        "The sequence presents plan selection, payment gateway interaction, transaction confirmation, and subscription activation.",
    ),
    Feature(
        "admin_manage_user",
        "Admin Manage User",
        "This feature allows administrators to search, inspect, activate, deactivate, or manage user accounts across platform roles.",
        "admin_manage_user_class_diagram_vertical_fit.png",
        "login-admin-manage-user.drawio.png",
        "Admin Manage User Sequence Diagram",
        "The sequence shows administrator authentication, user list retrieval, account action submission, update processing, and result feedback.",
    ),
    Feature(
        "admin_manage_job_content",
        "Admin Manage Job Content",
        "This feature supports administrative moderation of job content to maintain posting quality and policy compliance.",
        "admin_manage_job_content_class_diagram_vertical_fit.png",
        "login-admin_manage_job_content.drawio.png",
        "Admin Manage Job Content Sequence Diagram",
        "The sequence describes job moderation list retrieval, content review, approval or rejection action, and status update.",
    ),
    Feature(
        "admin_dashboard_statistics",
        "Admin Dashboard Statistics",
        "This feature provides administrators with platform metrics such as users, jobs, applications, subscriptions, and operational activity.",
        "admin_dashboard_statistics_class_diagram_vertical_fit.png",
        "login-admin_dashboard_statistics.drawio.png",
        "Admin Dashboard Statistics Sequence Diagram",
        "The sequence shows dashboard request handling, metric aggregation from services and repositories, and chart data response.",
    ),
]


@dataclass
class ClassSpec:
    name: str
    stereotypes: set[str] = field(default_factory=set)
    attributes: set[str] = field(default_factory=set)
    methods: set[str] = field(default_factory=set)
    sources: set[str] = field(default_factory=set)


def set_cell_text(cell, text: str, bold: bool = False, font_size: int = 9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_widths(table, widths: list[float]):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def clean_blocks_from_heading(doc: Document, heading_text: str):
    body = doc.element.body
    blocks = list(body)
    start = None
    for i, block in enumerate(blocks):
        text = "".join(node.text or "" for node in block.iter() if node.tag == qn("w:t"))
        if text.strip().startswith(heading_text):
            start = i
            break
    if start is None:
        raise ValueError(f"Could not find heading: {heading_text}")
    sect_pr = body.sectPr
    for block in blocks[start:]:
        if block is sect_pr:
            continue
        body.remove(block)


def fit_image_size(image_path: Path, max_width=6.4, max_height=7.25):
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return Inches(width), Inches(height)


def add_body_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(10.5)
    return p


def add_figure(doc: Document, image_path: Path, figure_no: int, caption: str, max_height=7.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width, height = fit_image_size(image_path, max_height=max_height)
    p.add_run().add_picture(str(image_path), width=width, height=height)

    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(6)
    run = caption_p.add_run(f"Figure {figure_no}. {caption}")
    run.bold = True
    run.font.size = Pt(10)


def normalize_member(member: str) -> str:
    return re.sub(r"\s+", " ", member.strip())


def parse_class_specs() -> dict[str, ClassSpec]:
    specs: dict[str, ClassSpec] = {}
    pattern = re.compile(r"^class\s+([A-Za-z_][\w]*)\s*\{(.*?)^\}", re.M | re.S)
    for path in sorted(CLASS_MMD_DIR.glob("*.mmd")):
        source = path.stem.replace("_class_diagram_vertical_fit", "")
        text = path.read_text(encoding="utf-8")
        for match in pattern.finditer(text):
            name = match.group(1)
            body = match.group(2)
            spec = specs.setdefault(name, ClassSpec(name=name))
            spec.sources.add(source)
            for raw_line in body.splitlines():
                line = raw_line.strip()
                if not line:
                    continue
                if line.startswith("<<") and line.endswith(">>"):
                    spec.stereotypes.add(line.strip("<>"))
                elif line[0] in "+-#~":
                    target = spec.methods if "(" in line and ")" in line else spec.attributes
                    target.add(normalize_member(line))
    return specs


def visibility(symbol: str) -> str:
    return {
        "+": "public",
        "-": "private",
        "#": "protected",
        "~": "package",
    }.get(symbol, "default")


def parse_attribute(member: str):
    vis = visibility(member[0])
    body = member[1:].strip()
    parts = body.split()
    if len(parts) >= 2:
        data_type = " ".join(parts[:-1])
        name = parts[-1]
    else:
        data_type = "unspecified"
        name = body
    return name, vis, data_type


def parse_method(member: str):
    vis = visibility(member[0])
    body = member[1:].strip()
    name_part, return_type = body, "void"
    after = re.match(r"(.+\))\s+(.+)$", body)
    if after:
        name_part, return_type = after.group(1), after.group(2)
    name = name_part.split("(", 1)[0].strip()
    params_raw = name_part.split("(", 1)[1].rsplit(")", 1)[0].strip() if "(" in name_part else ""
    params = []
    if params_raw:
        for index, param in enumerate([p.strip() for p in params_raw.split(",") if p.strip()], start=1):
            pieces = param.split()
            if len(pieces) == 1:
                params.append((f"param{index}", pieces[0]))
            else:
                params.append((pieces[-1], " ".join(pieces[:-1])))
    return name, vis, return_type, params


def class_package(spec: ClassSpec) -> str:
    stereo = " ".join(sorted(spec.stereotypes)).lower()
    name = spec.name.lower()
    if "react" in stereo or "frontend" in stereo or name.endswith("page") or name.endswith("storage"):
        return "frontend"
    if "controller" in stereo or spec.name.endswith("Controller"):
        return "controller"
    if "repository" in stereo or spec.name.endswith("Repository") or "database table" in stereo or spec.name.endswith("Table"):
        return "repository_database"
    if "dto" in stereo or "request" in name or "response" in name or "record" in stereo:
        return "dto_record"
    if "entity" in stereo or "domain" in stereo:
        return "domain_entity"
    if "external" in stereo or "gateway" in stereo or "component" in stereo or "spring" in stereo or "exception" in stereo:
        return "support_external"
    if "service" in stereo or spec.name.endswith("Service"):
        return "service"
    return "support_external"


def add_class_table(doc: Document, spec: ClassSpec):
    add_body_paragraph(
        doc,
        f"{spec.name} is a {', '.join(sorted(spec.stereotypes)) or 'design class'} used in "
        f"{', '.join(sorted(spec.sources))}.",
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, [0.55, 1.7, 4.05])
    for cell, text in zip(table.rows[0].cells, ["No", "Name", "Description"]):
        set_cell_text(cell, text, bold=True)
        shade_cell(cell, "D9EAF7")

    row = table.add_row()
    row.cells[0].merge(row.cells[2])
    set_cell_text(row.cells[0], "Attributes", bold=True)
    shade_cell(row.cells[0], "F2F2F2")

    attrs = sorted(spec.attributes)
    if attrs:
        for index, member in enumerate(attrs, start=1):
            name, vis, data_type = parse_attribute(member)
            row = table.add_row()
            set_cell_text(row.cells[0], f"{index:02d}")
            set_cell_text(row.cells[1], name)
            set_cell_text(
                row.cells[2],
                f"Visibility: {vis}\nType: {data_type}\nPurpose: Stores data required by {spec.name} for its design responsibility.",
            )
    else:
        row = table.add_row()
        set_cell_text(row.cells[0], "01")
        set_cell_text(row.cells[1], "N/A")
        set_cell_text(row.cells[2], "No explicit attributes are defined in the class diagram.")

    row = table.add_row()
    row.cells[0].merge(row.cells[2])
    set_cell_text(row.cells[0], "Methods/Operations", bold=True)
    shade_cell(row.cells[0], "F2F2F2")

    methods = sorted(spec.methods)
    if methods:
        for index, member in enumerate(methods, start=1):
            name, vis, return_type, params = parse_method(member)
            row = table.add_row()
            set_cell_text(row.cells[0], f"{index:02d}")
            set_cell_text(row.cells[1], name)
            param_lines = ["Parameters:"]
            if params:
                for param_name, param_type in params:
                    param_lines.append(f"- {param_name}: {param_type}, input data used by this operation.")
            else:
                param_lines.append("- None")
            set_cell_text(
                row.cells[2],
                f"Visibility: {vis}\nReturn: {return_type}\nPurpose: Performs the {name} responsibility for {spec.name}.\n"
                + "\n".join(param_lines),
            )
    else:
        row = table.add_row()
        set_cell_text(row.cells[0], "01")
        set_cell_text(row.cells[1], "N/A")
        set_cell_text(row.cells[2], "No explicit operations are defined in the class diagram.")

    doc.add_paragraph()


def add_detailed_design(doc: Document):
    doc.add_heading("2. Detailed Design", level=2)
    add_body_paragraph(
        doc,
        "This section provides the detailed design for the major system functions identified in the SRS. "
        "Each feature includes the class diagram and the corresponding sequence diagram that describes the runtime interaction flow.",
    )
    figure_no = 1
    for idx, feature in enumerate(FEATURES, start=1):
        doc.add_page_break()
        doc.add_heading(f"2.{idx} {feature.title}", level=3)
        add_body_paragraph(doc, feature.description)

        doc.add_heading(f"2.{idx}.1 Class Diagram", level=4)
        add_body_paragraph(
            doc,
            "This part presents the class structure, core responsibilities, and relationships used to implement the feature.",
        )
        add_figure(doc, CLASS_IMAGE_DIR / feature.class_image, figure_no, f"{feature.title} Class Diagram")
        figure_no += 1

        doc.add_heading(f"2.{idx}.2 {feature.sequence_title}", level=4)
        add_body_paragraph(doc, feature.sequence_description)
        if feature.sequence_image:
            add_figure(doc, SEQUENCE_IMAGE_DIR / feature.sequence_image, figure_no, feature.sequence_title, max_height=7.05)
            figure_no += 1
        else:
            add_body_paragraph(doc, feature.sequence_description)
    return figure_no


def add_class_specifications(doc: Document):
    specs = parse_class_specs()
    packages = defaultdict(list)
    for spec in specs.values():
        packages[class_package(spec)].append(spec)

    package_titles = [
        ("frontend", "frontend"),
        ("controller", "controller"),
        ("service", "service"),
        ("domain_entity", "domain_entity"),
        ("dto_record", "dto_record"),
        ("repository_database", "repository_database"),
        ("support_external", "support_external"),
    ]

    doc.add_page_break()
    doc.add_heading("3. Class Specifications", level=2)
    add_body_paragraph(
        doc,
        "This section provides detailed specifications for the classes extracted from the Mermaid class diagrams. "
        "Duplicate classes used by multiple features are consolidated and referenced once.",
    )

    section_no = 1
    for package_key, package_name in package_titles:
        package_specs = sorted(packages.get(package_key, []), key=lambda item: item.name)
        if not package_specs:
            continue
        doc.add_heading(f"3.{section_no} {package_name}", level=3)
        add_body_paragraph(
            doc,
            f"This section provides the detailed specifications for the classes in the package {package_name}.",
        )
        for class_no, spec in enumerate(package_specs, start=1):
            doc.add_heading(f"3.{section_no}.{class_no} {spec.name}", level=4)
            add_class_table(doc, spec)
        section_no += 1


def add_other_specs(doc: Document):
    doc.add_page_break()
    doc.add_heading("4. Other Design Specifications", level=2)
    add_body_paragraph(
        doc,
        "No additional design specifications are required beyond the high-level design, detailed feature design, and class specifications documented above.",
    )


def main():
    copy2(SOURCE_DOCX, OUTPUT_DOCX)
    doc = Document(OUTPUT_DOCX)
    clean_blocks_from_heading(doc, "2. Detailed Design")
    add_detailed_design(doc)
    add_class_specifications(doc)
    add_other_specs(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
